"""NJ DLGS "Abstract of Ratables" parsing: authoritative tax-levy split.

The Division of Local Government Services publishes a per-county "Abstract of
Ratables" xlsx each year with, for every municipality, the levy broken into
school / municipal / county buckets (Section 12 of the abstract). This is the
authoritative source for "where do my property taxes go" — distinct from the
HLS scrape (`taxes.*.parquet`), which sums per-parcel billed/paid and also
picks up PILOTs, added assessments, and water/sewer.

Column layout (verified stable across Hudson 2021-2025):
  col24        Net County Taxes Apportioned Less State Aid (12A5)
  col25-27     Net County Library / Health / Open Space (12B a/b/c)
  col28-30     District / Reg-Consol / Local School (12Ci a/b/c)
  col31-33     Municipal Budget / Open Space / Library (12Cii a/b/c)
  col34        Total Levy on Which Tax Rate Is Computed (12D) -- checksum

We validate school+municipal+county == total for every row and raise on drift,
so a future year with a shifted layout fails loudly rather than emitting garbage.
"""
import json
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import urlopen

import click
import openpyxl
from utz import err

from .paths import DATA

ABSTRACTS_DIR = DATA / "dlgs" / "abstracts"
# DLGS only hosts xlsx abstracts for 2021+ (older years are statewide .xls).
ABSTRACT_URL = "https://www.nj.gov/treasury/taxation/lpt/absractratables/Hudson{year}.xlsx"
DEFAULT_OUT = DATA.parent / "www" / "public" / "data" / "jc_levy_split.json"

# Section-12 column indices (1-based) in the abstract sheet.
COL_COUNTY = [24, 25, 26, 27]     # 12A5 net-less-aid + 12B library/health/open-space
COL_SCHOOL = [28, 29, 30]         # 12Ci district/reg-consol/local
COL_MUNI = [31, 32, 33]           # 12Cii budget/open-space/library
COL_MUNI_BUDGET = 31              # 12Cii(A) alone -- the "$385M" municipal-budget figure
COL_TOTAL = 34                    # 12D checksum
JC_MCODE = "0906"


def _abstract_path(year: int) -> Path:
    return ABSTRACTS_DIR / f"Hudson{year}.xlsx"


def download_abstract(year: int, force: bool = False) -> Path:
    """Download the Hudson abstract xlsx for `year` if not already cached."""
    path = _abstract_path(year)
    if path.exists() and not force:
        return path
    url = ABSTRACT_URL.format(year=year)
    err(f"Downloading {url}")
    try:
        with urlopen(url, timeout=60) as resp:
            content = resp.read()
    except HTTPError as e:
        if e.code == 404:
            raise FileNotFoundError(f"No abstract hosted for {year} ({url})") from e
        raise
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    return path


def parse_abstract(year: int) -> list[dict]:
    """Parse one Hudson abstract into per-municipality levy splits."""
    ws = openpyxl.load_workbook(_abstract_path(year), data_only=True).active
    rows: list[dict] = []
    for r in range(6, ws.max_row + 1):
        mcode = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        if not mcode or not name:
            continue
        mcode = str(mcode).strip()
        if not mcode.startswith("09"):  # Hudson muni codes are 09xx
            continue

        def s(cols) -> float:
            return sum(ws.cell(r, c).value or 0 for c in cols)

        school = s(COL_SCHOOL)
        municipal = s(COL_MUNI)
        county = s(COL_COUNTY)
        total = ws.cell(r, COL_TOTAL).value or 0
        if abs((school + municipal + county) - total) > 1:
            raise ValueError(
                f"{year} {name}: school+muni+county ({school + municipal + county:.0f}) "
                f"!= total levy ({total:.0f}); abstract column layout may have shifted",
            )
        rows.append({
            "year": year,
            "mcode": mcode,
            "name": str(name).strip(),
            "school": round(school),
            "municipal": round(municipal),
            "municipal_budget": round(ws.cell(r, COL_MUNI_BUDGET).value or 0),
            "county": round(county),
            "total_levy": round(total),
        })
    return rows


@click.group()
def dlgs():
    """NJ DLGS Abstract of Ratables (authoritative tax-levy split)."""


@dlgs.command()
@click.option("-a", "--all-munis", is_flag=True, help="Emit all Hudson munis (default: Jersey City only)")
@click.option("-f", "--force-download", is_flag=True, help="Re-download abstracts even if cached")
@click.option("-o", "--output", default=str(DEFAULT_OUT), show_default=True, help="Output JSON path")
@click.option("-y", "--years", default="2021-2025", show_default=True, help="Year range (e.g. '2021-2025')")
def levy(all_munis: bool, force_download: bool, output: str, years: str):
    """Extract the school/municipal/county levy split over time -> JSON."""
    start, end = (int(x) for x in years.split("-")) if "-" in years else (int(years), int(years))
    year_list = list(range(start, end + 1))

    all_rows: list[dict] = []
    for year in year_list:
        download_abstract(year, force=force_download)
        rows = parse_abstract(year)
        if not all_munis:
            rows = [r for r in rows if r["mcode"] == JC_MCODE]
        all_rows.extend(rows)
        err(f"{year}: {len(rows)} muni row(s)")

    for r in all_rows:
        if not all_munis:
            r.pop("mcode", None)
            r.pop("name", None)

    out_path = Path(output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(all_rows, indent=2) + "\n")
    err(f"Wrote {len(all_rows)} rows to {output}")


if __name__ == "__main__":
    dlgs()
